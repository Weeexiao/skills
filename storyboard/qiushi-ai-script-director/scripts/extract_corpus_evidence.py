#!/usr/bin/env python3
"""Build a portable evidence index from the source lapping-script corpus.

This is a maintainer tool.  It reads a supplied corpus directory, converts legacy
``.xls`` files to temporary ``.xlsx`` copies with LibreOffice, and emits JSONL and
Markdown reports.  The generated index deliberately stores only portable source
labels (case id, relative file name, sheet and row) rather than a creator's
absolute computer path.
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import sys
import tempfile
from collections import Counter, defaultdict
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


TEXT_LIMIT = 1600


def clean(value: object) -> str:
    if value is None:
        return ""
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text[:TEXT_LIMIT]


def source_files(corpus: Path) -> list[Path]:
    return sorted(
        p
        for p in corpus.rglob("*")
        if p.is_file() and p.suffix.lower() in {".xls", ".xlsx"} and not p.name.startswith("~$")
    )


def as_xlsx(source: Path, temp_root: Path, soffice: str) -> Path:
    if source.suffix.lower() == ".xlsx":
        return source
    target_dir = temp_root / source.parent.name
    target_dir.mkdir(parents=True, exist_ok=True)
    result = subprocess.run(
        [soffice, "--headless", "--convert-to", "xlsx", "--outdir", str(target_dir), str(source)],
        capture_output=True,
        text=True,
        timeout=90,
    )
    converted = target_dir / f"{source.stem}.xlsx"
    if result.returncode != 0 or not converted.exists():
        raise RuntimeError((result.stderr or result.stdout or "conversion failed").strip())
    return converted


def iter_sheet_rows(workbook_path: Path) -> Iterable[tuple[str, int, list[str]]]:
    book = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        for sheet in book.worksheets:
            blank_run = 0
            for row_number, cells in enumerate(sheet.iter_rows(values_only=True), start=1):
                values = [clean(cell) for cell in cells]
                if not any(values):
                    blank_run += 1
                    if blank_run >= 80:
                        break
                    continue
                blank_run = 0
                yield sheet.title, row_number, values
    finally:
        book.close()


def relevant_text(values: list[str]) -> str:
    return " | ".join(value for value in values if value)[:TEXT_LIMIT]


def main() -> int:
    parser = argparse.ArgumentParser(description="Extract a portable lapping-script corpus evidence index.")
    parser.add_argument("corpus", type=Path, help="Directory containing unique source Excel lapping scripts")
    parser.add_argument("output", type=Path, help="Directory for evidence-index.jsonl and corpus-coverage.md")
    parser.add_argument("--soffice", default="soffice", help="LibreOffice/soffice executable")
    args = parser.parse_args()

    corpus = args.corpus.resolve()
    out = args.output.resolve()
    out.mkdir(parents=True, exist_ok=True)
    files = source_files(corpus)
    if not files:
        raise SystemExit(f"No .xls/.xlsx files found in: {corpus}")

    topic_counts: Counter[str] = Counter()
    failures: list[dict[str, str]] = []
    summaries: list[dict[str, object]] = []
    row_records: list[dict[str, object]] = []
    with tempfile.TemporaryDirectory(prefix="qiushi-corpus-") as temp_dir_string:
        temp_dir = Path(temp_dir_string)
        for index, source in enumerate(files, start=1):
            rel = source.relative_to(corpus)
            topic = rel.parts[0] if len(rel.parts) > 1 else "未分类"
            case_id = f"QS-{index:03d}"
            try:
                workbook = as_xlsx(source, temp_dir, args.soffice)
                rows = list(iter_sheet_rows(workbook))
                sheet_names = sorted({sheet for sheet, _, _ in rows})
                text_rows = []
                for sheet, row_number, values in rows:
                    text = relevant_text(values)
                    if text:
                        record = {
                            "case_id": case_id,
                            "topic": topic,
                            "source_file": source.name,
                            "source_relative": str(rel),
                            "sheet": sheet,
                            "row": row_number,
                            "text": text,
                        }
                        row_records.append(record)
                        text_rows.append(text)
                summaries.append(
                    {
                        "case_id": case_id,
                        "topic": topic,
                        "source_file": source.name,
                        "source_relative": str(rel),
                        "sheets": sheet_names,
                        "nonempty_rows": len(text_rows),
                        "preview": " / ".join(text_rows[:3])[:1000],
                    }
                )
                topic_counts[topic] += 1
            except Exception as exc:  # keep the coverage report honest
                failures.append({"case_id": case_id, "source_relative": str(rel), "error": str(exc)[:500]})

    jsonl_path = out / "evidence-index.jsonl"
    with jsonl_path.open("w", encoding="utf-8") as handle:
        for record in row_records:
            handle.write(json.dumps(record, ensure_ascii=False) + "\n")

    summary_path = out / "case-summary.json"
    summary_path.write_text(json.dumps(summaries, ensure_ascii=False, indent=2), encoding="utf-8")

    coverage = [
        "# 秋拾拉片案例库：提炼覆盖报告",
        "",
        f"- 主样本目录：`{corpus.name}`（报告中不写入卖家绝对路径）",
        f"- 发现表格：{len(files)} 份",
        f"- 成功读取：{len(summaries)} 份",
        f"- 可检索非空行：{len(row_records)} 行",
        f"- 读取失败：{len(failures)} 份",
        "",
        "## 题材覆盖",
        "",
    ]
    for topic, count in topic_counts.most_common():
        coverage.append(f"- {topic}：{count} 份")
    coverage.extend(["", "## 案例清单", ""])
    for item in summaries:
        coverage.append(
            f"- {item['case_id']}｜{item['topic']}｜{item['source_file']}｜"
            f"工作表：{', '.join(item['sheets']) or '未命名'}｜非空行：{item['nonempty_rows']}"
        )
    if failures:
        coverage.extend(["", "## 未成功读取（保留为后续人工核对，不计入方法证据）", ""])
        for failed in failures:
            coverage.append(f"- {failed['case_id']}｜{failed['source_relative']}｜{failed['error']}")
    (out / "corpus-coverage.md").write_text("\n".join(coverage) + "\n", encoding="utf-8")

    print(json.dumps({"files": len(files), "read": len(summaries), "rows": len(row_records), "failed": len(failures)}, ensure_ascii=False))
    return 0 if not failures else 2


if __name__ == "__main__":
    sys.exit(main())
